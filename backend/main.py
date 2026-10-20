import io
import json
import os

import numpy as np
import pandas as pd
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from groq import Groq
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

load_dotenv()

app = FastAPI(title="InsightAI - AI Data Analyst")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
    "http://localhost:5173",
    "http://localhost:5174",
    "http://localhost:5175",
    "http://localhost:5176",
    "http://localhost:5177",
    "http://localhost:5178",
    "https://insightai-data-analyst.vercel.app",
],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")

# Uploaded datasets are stored in memory.
# Upload the CSV again after restarting the backend.
DATASETS = {}


# --------------------------------------------------
# GENERAL HELPERS
# --------------------------------------------------

def get_groq_client():
    api_key = os.getenv("GROQ_API_KEY")

    if not api_key:
        return None

    return Groq(api_key=api_key)


def clean_value(value):
    """Convert NumPy/Pandas values into JSON-safe values."""
    if value is None:
        return None

    if isinstance(value, (np.integer,)):
        return int(value)

    if isinstance(value, (np.floating,)):
        if not np.isfinite(value):
            return None
        return float(value)

    if isinstance(value, (np.bool_,)):
        return bool(value)

    if pd.isna(value):
        return None

    if isinstance(value, pd.Timestamp):
        return value.isoformat()

    return value


def dataframe_records(df):
    """Convert a DataFrame into JSON-safe records."""
    records = (
        df.astype(object)
        .where(pd.notna(df), None)
        .to_dict(orient="records")
    )

    return [
        {
            str(key): clean_value(value)
            for key, value in row.items()
        }
        for row in records
    ]


def profile(df):
    """Return basic dataset and column information."""
    columns = []

    for column in df.columns:
        series = df[column]

        columns.append({
            "name": str(column),
            "dtype": str(series.dtype),
            "missing": int(series.isna().sum()),
            "unique": int(series.nunique(dropna=True)),
            "sample": [
                clean_value(value)
                for value in series.dropna().head(5).tolist()
            ],
        })

    return {
        "rows": int(df.shape[0]),
        "columns": int(df.shape[1]),
        "missing_values": int(df.isna().sum().sum()),
        "duplicate_rows": int(df.duplicated().sum()),
        "column_profiles": columns,
    }


def data_quality_summary(df):
    missing_by_column = {
        str(column): int(count)
        for column, count in df.isna().sum().items()
        if count
    }

    outliers_by_column = {}
    for column in df.select_dtypes(include=np.number).columns:
        values = df[column].dropna()
        if values.empty:
            continue

        lower_quartile = values.quantile(0.25)
        upper_quartile = values.quantile(0.75)
        spread = upper_quartile - lower_quartile
        if spread <= 0:
            continue

        count = int(
            (
                (values < lower_quartile - 1.5 * spread)
                | (values > upper_quartile + 1.5 * spread)
            ).sum()
        )
        if count:
            outliers_by_column[str(column)] = count

    type_issues = {}
    inconsistent_categories = {}
    for column in df.columns:
        dtype = df[column].dtype
        if not (
            pd.api.types.is_object_dtype(dtype)
            or isinstance(dtype, pd.CategoricalDtype)
            or pd.api.types.is_string_dtype(dtype)
        ):
            continue
        values = df[column].dropna().astype(str)
        if values.empty:
            continue

        numeric_values = pd.to_numeric(values, errors="coerce")
        invalid_numeric_count = int(numeric_values.isna().sum())
        if numeric_values.notna().any() and invalid_numeric_count:
            type_issues[str(column)] = invalid_numeric_count

        normalized = values.str.strip().str.casefold()
        variations = int(values.nunique() - normalized.nunique())
        if variations:
            inconsistent_categories[str(column)] = variations

    return {
        "missing_values": {
            "count": int(df.isna().sum().sum()),
            "columns": missing_by_column,
        },
        "duplicate_rows": int(df.duplicated().sum()),
        "outliers": {
            "count": int(sum(outliers_by_column.values())),
            "columns": outliers_by_column,
        },
        "type_issues": {
            "count": int(sum(type_issues.values())),
            "columns": type_issues,
        },
        "inconsistent_categories": {
            "count": int(sum(inconsistent_categories.values())),
            "columns": inconsistent_categories,
        },
    }


def clean_dataframe(source, operations):
    if not isinstance(operations, dict):
        raise HTTPException(status_code=400, detail="Cleaning options must be an object.")

    df = source.copy(deep=True)
    initial_rows = len(df)
    remove_duplicates = operations.get("remove_duplicates", False)
    if not isinstance(remove_duplicates, bool):
        raise HTTPException(
            status_code=400,
            detail="remove_duplicates must be true or false.",
        )

    def validate_columns(values, option_name):
        if not isinstance(values, list) or any(
            not isinstance(column, str) or column not in df.columns
            for column in values
        ):
            raise HTTPException(
                status_code=400,
                detail=f"{option_name} contains an unknown column.",
            )

    drop_columns = operations.get("drop_columns", [])
    validate_columns(drop_columns, "drop_columns")
    if len(set(drop_columns)) >= len(df.columns) and drop_columns:
        raise HTTPException(
            status_code=400,
            detail="Cleaning cannot remove every column.",
        )
    if drop_columns:
        df = df.drop(columns=drop_columns)

    convert_types = operations.get("convert_types", {})
    if not isinstance(convert_types, dict):
        raise HTTPException(status_code=400, detail="convert_types must be an object.")
    for column, target_type in convert_types.items():
        if column not in df.columns:
            if column in drop_columns:
                continue
            raise HTTPException(status_code=400, detail=f"Unknown column: {column}.")
        try:
            if target_type == "text":
                df[column] = df[column].astype("string")
            elif target_type == "numeric":
                converted = pd.to_numeric(df[column], errors="coerce")
                invalid = df[column].notna() & converted.isna()
                if invalid.any():
                    raise ValueError(
                        f"{int(invalid.sum())} value(s) in {column} are not numeric."
                    )
                df[column] = converted
            elif target_type == "integer":
                converted = pd.to_numeric(df[column], errors="coerce")
                invalid = df[column].notna() & (
                    converted.isna() | (converted % 1 != 0)
                )
                if invalid.any():
                    raise ValueError(
                        f"{int(invalid.sum())} value(s) in {column} are not integers."
                    )
                df[column] = converted.astype("Int64")
            elif target_type == "datetime":
                converted = pd.to_datetime(df[column], errors="coerce")
                invalid = df[column].notna() & converted.isna()
                if invalid.any():
                    raise ValueError(
                        f"{int(invalid.sum())} value(s) in {column} are not valid dates."
                    )
                df[column] = converted
            elif target_type == "category":
                df[column] = df[column].astype("category")
            else:
                raise HTTPException(
                    status_code=400,
                    detail=f"Unsupported target type for {column}.",
                )
        except (TypeError, OverflowError) as error:
            raise HTTPException(
                status_code=400,
                detail=f"Could not convert {column}: {error}",
            ) from error
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error

    fill_missing = operations.get("fill_missing", {})
    if not isinstance(fill_missing, dict):
        raise HTTPException(status_code=400, detail="fill_missing must be an object.")
    for column, option in fill_missing.items():
        if column not in df.columns:
            if column in drop_columns:
                continue
            raise HTTPException(status_code=400, detail=f"Unknown column: {column}.")
        if not isinstance(option, dict):
            raise HTTPException(
                status_code=400,
                detail=f"Missing-value options for {column} must be an object.",
            )
        strategy = option.get("strategy", "keep")
        if strategy == "keep":
            continue
        if strategy == "drop_rows":
            df = df.dropna(subset=[column])
        elif strategy in ("mean", "median", "mode"):
            if strategy in ("mean", "median") and not pd.api.types.is_numeric_dtype(
                df[column]
            ):
                raise HTTPException(
                    status_code=400,
                    detail=f"{strategy.title()} filling requires a numeric column: {column}.",
                )
            if strategy == "mean":
                fill_value = df[column].mean()
            elif strategy == "median":
                fill_value = df[column].median()
            else:
                modes = df[column].mode(dropna=True)
                if modes.empty:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Cannot find a mode for empty column {column}.",
                    )
                fill_value = modes.iloc[0]
            df[column] = df[column].fillna(fill_value)
        elif strategy == "constant":
            fill_value = option.get("value", "")
            if not isinstance(fill_value, (str, int, float, bool)):
                raise HTTPException(
                    status_code=400,
                    detail=f"Fill value for {column} must be text or numeric.",
                )
            if isinstance(df[column].dtype, pd.CategoricalDtype):
                if fill_value not in df[column].cat.categories:
                    df[column] = df[column].cat.add_categories([fill_value])
            try:
                df[column] = df[column].fillna(fill_value)
            except (TypeError, ValueError, OverflowError) as error:
                raise HTTPException(
                    status_code=400,
                    detail=f"Could not fill missing values in {column}: {error}",
                ) from error
        else:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported missing-value strategy for {column}.",
            )

    if remove_duplicates:
        df = df.drop_duplicates()

    outliers = operations.get("outliers", {})
    if not isinstance(outliers, dict):
        raise HTTPException(status_code=400, detail="outliers must be an object.")
    for column, strategy in outliers.items():
        if column not in df.columns:
            if column in drop_columns:
                continue
            raise HTTPException(status_code=400, detail=f"Unknown column: {column}.")
        if not pd.api.types.is_numeric_dtype(df[column]):
            raise HTTPException(
                status_code=400,
                detail=f"Outlier handling requires a numeric column: {column}.",
            )
        if strategy == "keep":
            continue
        values = df[column].dropna()
        if values.empty:
            continue
        lower_quartile = values.quantile(0.25)
        upper_quartile = values.quantile(0.75)
        spread = upper_quartile - lower_quartile
        lower_bound = lower_quartile - 1.5 * spread
        upper_bound = upper_quartile + 1.5 * spread
        outlier_mask = (df[column] < lower_bound) | (df[column] > upper_bound)
        if strategy == "cap":
            df[column] = df[column].clip(lower_bound, upper_bound)
        elif strategy == "remove":
            df = df.loc[~outlier_mask].copy()
        else:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported outlier strategy for {column}.",
            )

    if df.empty:
        raise HTTPException(
            status_code=400,
            detail="These cleaning options would leave no data rows.",
        )

    return df, {
        "rows_before": initial_rows,
        "rows_after": len(df),
        "rows_removed": initial_rows - len(df),
    }


def build_dataset_context(df):
    """Build a concise description for the AI model."""
    description = {
        "rows": int(df.shape[0]),
        "columns": list(map(str, df.columns)),
        "dtypes": {
            str(column): str(df[column].dtype)
            for column in df.columns
        },
        "missing_values": int(df.isna().sum().sum()),
        "duplicate_rows": int(df.duplicated().sum()),
        "sample_data": dataframe_records(df.head(10)),
    }

    numeric = df.select_dtypes(include=np.number)

    if not numeric.empty:
        description["numeric_summary"] = json.loads(
            numeric.describe().round(2).to_json()
        )

    categorical = df.select_dtypes(
        include=["object", "category", "bool"]
    )

    description["categorical_values"] = {}

    for column in categorical.columns:
        description["categorical_values"][str(column)] = (
            df[column]
            .value_counts(dropna=False)
            .head(15)
            .to_dict()
        )

    return json.dumps(
        description,
        default=str,
        ensure_ascii=False,
    )


def find_column(df, text):
    """Find a column by case-insensitive exact name or substring."""
    text = text.lower().strip()

    for column in df.columns:
        if str(column).lower() == text:
            return column

    for column in df.columns:
        if str(column).lower() in text:
            return column

    return None


def find_mentioned_column(df, question, columns=None):
    """Find the longest column name mentioned in a question."""
    question_lower = question.lower()

    if columns is None:
        columns = list(df.columns)

    matches = []

    for column in columns:
        name = str(column).lower()

        if name and name in question_lower:
            matches.append(column)

    if not matches:
        return None

    return max(matches, key=lambda item: len(str(item)))


def fmt(value):
    """Format a number for readable answers."""
    if value is None:
        return "N/A"

    try:
        number = float(value)

        if not np.isfinite(number):
            return "N/A"

        if number.is_integer():
            return f"{int(number):,}"

        return f"{number:,.2f}"

    except (TypeError, ValueError):
        return str(value)


def contains_any(text, phrases):
    return any(phrase in text for phrase in phrases)


def group_summary(df, group_column, value_column, operation="sum"):
    """Summarize a numeric column by category."""
    working = df[[group_column, value_column]].copy()

    working[value_column] = pd.to_numeric(
        working[value_column],
        errors="coerce",
    )

    working = working.dropna(subset=[value_column])

    if working.empty:
        return pd.DataFrame(columns=[group_column, value_column])

    grouped = working.groupby(
        group_column,
        dropna=False,
    )[value_column]

    if operation == "mean":
        result = grouped.mean()
    elif operation == "count":
        result = grouped.count()
    elif operation == "max":
        result = grouped.max()
    elif operation == "min":
        result = grouped.min()
    else:
        result = grouped.sum()

    return result.sort_values(
        ascending=False
    ).reset_index()


def sort_chart_data(summary, category_column):
    """Sort date-like categories chronologically."""
    if summary.empty or category_column not in summary.columns:
        return summary

    category_name = str(category_column).lower()

    name_suggests_date = any(
        word in category_name
        for word in ["date", "datetime", "timestamp", "time"]
    )

    values = summary[category_column]

    parsed_dates = pd.to_datetime(
        values,
        errors="coerce",
    )

    valid_ratio = (
        float(parsed_dates.notna().mean())
        if len(parsed_dates) > 0
        else 0.0
    )

    if name_suggests_date or valid_ratio >= 0.8:
        if valid_ratio >= 0.8:
            result = summary.copy()
            result["_chart_date_sort"] = parsed_dates

            result = result.sort_values(
                "_chart_date_sort",
                ascending=True,
                na_position="last",
            )

            return result.drop(
                columns=["_chart_date_sort"]
            ).reset_index(drop=True)

    return summary.reset_index(drop=True)


def build_histogram(series, max_bins=10):
    """Create a numeric histogram with bin midpoints and frequencies."""
    values = pd.to_numeric(
        series,
        errors="coerce",
    ).dropna().to_numpy(dtype=float)

    if len(values) == 0:
        return []

    bin_count = min(
        max_bins,
        max(1, int(np.ceil(np.sqrt(len(values))))),
    )

    counts, edges = np.histogram(
        values,
        bins=bin_count,
    )

    histogram = []

    for index, count in enumerate(counts):
        midpoint = (edges[index] + edges[index + 1]) / 2

        histogram.append({
            "value": round(float(midpoint), 2),
            "count": int(count),
        })

    return histogram


def wants_group_comparison(question):
    question = question.lower()

    return contains_any(question, [
        "compare",
        "comparison",
        "by region",
        "by product",
        "breakdown",
        "group by",
        "grouped by",
        "each region",
        "each product",
        "per region",
        "per product",
        "across regions",
        "across products",
        "sales for every",
        "total sales for",
    ])


def detect_operation(question):
    question = question.lower()

    if contains_any(question, ["average", "mean", "avg"]):
        return "mean"

    if contains_any(question, ["count", "number of records"]):
        return "count"

    if contains_any(question, ["maximum", "max value"]):
        return "max"

    if contains_any(question, ["minimum", "min value"]):
        return "min"

    return "sum"


# --------------------------------------------------
# BUSINESS INSIGHTS
# --------------------------------------------------

def build_business_insights(df):
    """
    Calculate business KPIs, product performance, regional performance,
    profit margin, and recommendations from the complete uploaded dataset.
    """

    numeric_columns = list(
        df.select_dtypes(include=np.number).columns
    )

    categorical_columns = list(
        df.select_dtypes(
            include=["object", "category", "bool"]
        ).columns
    )

    # Identify sales or revenue column.
    sales_column = next(
        (
            column for column in numeric_columns
            if any(
                word in str(column).lower()
                for word in ["sales", "revenue", "turnover"]
            )
        ),
        numeric_columns[0] if numeric_columns else None,
    )

    # Identify profit column.
    profit_column = next(
        (
            column for column in numeric_columns
            if any(
                word in str(column).lower()
                for word in ["profit", "earnings"]
            )
        ),
        None,
    )

    # Identify product and region columns.
    product_column = next(
        (
            column for column in categorical_columns
            if any(
                word in str(column).lower()
                for word in ["product", "item", "service"]
            )
        ),
        None,
    )

    region_column = next(
        (
            column for column in categorical_columns
            if any(
                word in str(column).lower()
                for word in [
                    "region",
                    "location",
                    "territory",
                    "area",
                ]
            )
        ),
        None,
    )

    insights = {
        "total_sales": None,
        "total_profit": None,
        "average_sales": None,
        "profit_margin": None,
        "highest_sale": None,
        "top_product": None,
        "lowest_product": None,
        "product_performance": [],
        "best_region": None,
        "region_performance": [],
        "recommendations": [],
    }

    # Sales statistics.
    if sales_column is not None:
        sales = pd.to_numeric(
            df[sales_column],
            errors="coerce",
        ).dropna()

        if not sales.empty:
            insights["total_sales"] = round(float(sales.sum()), 2)
            insights["average_sales"] = round(float(sales.mean()), 2)
            insights["highest_sale"] = round(float(sales.max()), 2)

    # Profit and profit margin.
    if profit_column is not None:
        profit = pd.to_numeric(
            df[profit_column],
            errors="coerce",
        ).dropna()

        if not profit.empty:
            insights["total_profit"] = round(float(profit.sum()), 2)

            total_sales = insights["total_sales"]

            if total_sales not in (None, 0):
                insights["profit_margin"] = round(
                    float(profit.sum()) / total_sales * 100,
                    2,
                )

    # Product-level sales.
    if product_column is not None and sales_column is not None:
        product_summary = group_summary(
            df,
            product_column,
            sales_column,
            operation="sum",
        )

        if not product_summary.empty:
            insights["product_performance"] = [
                {
                    "name": str(row[product_column]),
                    "sales": round(float(row[sales_column]), 2),
                }
                for _, row in product_summary.iterrows()
            ]

            top = product_summary.iloc[0]
            lowest = product_summary.iloc[-1]

            insights["top_product"] = {
                "name": str(top[product_column]),
                "sales": round(float(top[sales_column]), 2),
            }

            insights["lowest_product"] = {
                "name": str(lowest[product_column]),
                "sales": round(float(lowest[sales_column]), 2),
            }

    # Region-level sales.
    if region_column is not None and sales_column is not None:
        region_summary = group_summary(
            df,
            region_column,
            sales_column,
            operation="sum",
        )

        if not region_summary.empty:
            insights["region_performance"] = [
                {
                    "name": str(row[region_column]),
                    "sales": round(float(row[sales_column]), 2),
                }
                for _, row in region_summary.iterrows()
            ]

            best = region_summary.iloc[0]

            insights["best_region"] = {
                "name": str(best[region_column]),
                "sales": round(float(best[sales_column]), 2),
            }

    # Recommendations based on the calculated results.
    top_product = insights["top_product"]
    lowest_product = insights["lowest_product"]
    best_region = insights["best_region"]
    margin = insights["profit_margin"]

    if top_product:
        insights["recommendations"].append(
            f"Prioritize analysis of {top_product['name']}, "
            f"the leading product with {fmt(top_product['sales'])} "
            "in sales. Investigate the factors behind its performance."
        )

    if lowest_product:
        insights["recommendations"].append(
            f"Review {lowest_product['name']}, which has the lowest "
            f"product sales at {fmt(lowest_product['sales'])}. "
            "Investigate demand, pricing, and customer preferences."
        )

    if best_region:
        insights["recommendations"].append(
            f"Investigate what drives sales in {best_region['name']}, "
            f"the leading region with {fmt(best_region['sales'])} "
            "in sales, and assess whether those practices can be repeated."
        )

    if margin is not None:
        insights["recommendations"].append(
            f"The overall profit margin is {margin:.2f}%. "
            "Compare profitability across products and regions to "
            "identify opportunities for improvement."
        )

    if not insights["recommendations"]:
        insights["recommendations"].append(
            "Explore numeric summaries and category breakdowns "
            "to identify useful trends in this dataset."
        )

    return insights


def build_advanced_insights(df):
    """Calculate additional verified business metrics from the dataset."""
    numeric_columns = list(
        df.select_dtypes(include=np.number).columns
    )
    categorical_columns = list(
        df.select_dtypes(
            include=["object", "category", "bool", "string"]
        ).columns
    )

    def find_matching_column(columns, keywords):
        return next(
            (
                column
                for column in columns
                if any(
                    keyword in str(column).lower()
                    for keyword in keywords
                )
            ),
            None,
        )

    sales_column = find_matching_column(
        numeric_columns,
        ["sales", "revenue", "turnover"],
    )
    quantity_column = find_matching_column(
        numeric_columns,
        ["quantity", "qty"],
    )
    category_column = find_matching_column(
        categorical_columns,
        ["category", "classification"],
    )
    region_column = find_matching_column(
        categorical_columns,
        ["region", "location", "territory", "area"],
    )
    product_column = find_matching_column(
        categorical_columns,
        ["product", "item", "service"],
    )

    total_sales = None
    top_sales_category = None
    best_region = None
    lowest_region = None
    top_2_products = None

    if sales_column is not None:
        sales = pd.to_numeric(df[sales_column], errors="coerce")
        sales_values = sales.dropna()
        total_sales = sales_values.sum() if not sales_values.empty else None

        if category_column is not None:
            category_summary = group_summary(
                df,
                category_column,
                sales_column,
                operation="sum",
            )
            if not category_summary.empty:
                top_category = category_summary.iloc[0]
                category_sales = clean_value(top_category[sales_column])
                top_sales_category = {
                    "category": clean_value(top_category[category_column]),
                    "total_sales": category_sales,
                    "percentage_of_total_sales": (
                        round(float(category_sales) / float(total_sales) * 100, 2)
                        if total_sales not in (None, 0)
                        else None
                    ),
                }

        if region_column is not None:
            region_summary = group_summary(
                df,
                region_column,
                sales_column,
                operation="sum",
            )
            if not region_summary.empty:
                best = region_summary.iloc[0]
                lowest = region_summary.iloc[-1]
                best_region = {
                    "region": clean_value(best[region_column]),
                    "sales": clean_value(best[sales_column]),
                }
                lowest_region = {
                    "region": clean_value(lowest[region_column]),
                    "sales": clean_value(lowest[sales_column]),
                }

        if product_column is not None:
            product_summary = group_summary(
                df,
                product_column,
                sales_column,
                operation="sum",
            ).head(2)
            if not product_summary.empty:
                products = [
                    {
                        "product": clean_value(row[product_column]),
                        "sales": clean_value(row[sales_column]),
                    }
                    for _, row in product_summary.iterrows()
                ]
                combined_sales = sum(
                    float(product["sales"] or 0)
                    for product in products
                )
                top_2_products = {
                    "products": products,
                    "combined_sales": clean_value(combined_sales),
                    "percentage_of_total_sales": (
                        round(combined_sales / float(total_sales) * 100, 2)
                        if total_sales not in (None, 0)
                        else None
                    ),
                }

    quantity = (
        pd.to_numeric(df[quantity_column], errors="coerce").dropna()
        if quantity_column is not None
        else pd.Series(dtype=float)
    )

    return {
        "top_sales_category": top_sales_category,
        "best_region": best_region,
        "lowest_region": lowest_region,
        "top_2_products": top_2_products,
        "total_quantity": (
            clean_value(quantity.sum()) if not quantity.empty else None
        ),
        "average_quantity": (
            round(float(quantity.mean()), 2) if not quantity.empty else None
        ),
        "highest_quantity": (
            clean_value(quantity.max()) if not quantity.empty else None
        ),
        "missing_cells": int(df.isna().sum().sum()),
        "duplicate_rows": int(df.duplicated().sum()),
    }


# --------------------------------------------------
# ROOT AND HEALTH
# --------------------------------------------------

@app.get("/")
def home():
    return {
        "message": "InsightAI AI Data Analyst API is running",
        "docs": "/docs",
        "health": "/api/health",
    }


@app.get("/api/health")
def health():
    return {
        "status": "healthy",
        "service": "InsightAI",
        "datasets_loaded": len(DATASETS),
        "groq_configured": bool(os.getenv("GROQ_API_KEY")),
    }


# --------------------------------------------------
# CSV UPLOAD
# --------------------------------------------------

@app.post("/api/upload")
async def upload_csv(file: UploadFile = File(...)):
    filename = file.filename or ""
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if extension not in ("csv", "xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only CSV and Excel files are supported.",
        )

    try:
        content = await file.read()

        # Limit uploaded files to 15 MB.
        if len(content) > 15 * 1024 * 1024:
            raise HTTPException(
                status_code=413,
                detail="The file is too large. Maximum size is 15 MB.",
            )

        if not content:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is empty.",
            )

        if extension == "xlsx":
            file.file.seek(0)
            df = pd.read_excel(file.file, engine="openpyxl")
        else:
            try:
                df = pd.read_csv(io.BytesIO(content))
            except UnicodeDecodeError:
                df = pd.read_csv(
                    io.BytesIO(content),
                    encoding="latin-1",
                )

        if df.empty:
            raise HTTPException(
                status_code=400,
                detail="The file contains no data rows.",
            )

        if len(df.columns) == 0:
            raise HTTPException(
                status_code=400,
                detail="No columns were found in the file.",
            )

        df.columns = [
            str(column).strip()
            for column in df.columns
        ]

        if any(not column for column in df.columns):
            raise HTTPException(
                status_code=400,
                detail="Every column must have a name.",
            )

        if df.columns.duplicated().any():
            raise HTTPException(
                status_code=400,
                detail=(
                    "The file contains duplicate column names. "
                    "Rename them and try again."
                ),
            )

        dataset_id = os.urandom(8).hex()

        DATASETS[dataset_id] = {
            "filename": filename,
            "df": df,
            "original_df": df.copy(deep=True),
        }

        numeric_columns = list(
            map(
                str,
                df.select_dtypes(include=np.number).columns,
            )
        )

        categorical_columns = list(
            map(
                str,
                df.select_dtypes(
                    include=["object", "category", "bool"]
                ).columns,
            )
        )

        return {
            "success": True,
            "dataset_id": dataset_id,
            "filename": filename,
            "message": "File uploaded successfully.",
            "preview": dataframe_records(df.head(10)),
            "columns": list(map(str, df.columns)),
            "numeric_columns": numeric_columns,
            "categorical_columns": categorical_columns,
            "profile": profile(df),
            "quality": data_quality_summary(df),
        }

    except HTTPException:
        raise

    except pd.errors.EmptyDataError:
        raise HTTPException(
            status_code=400,
            detail="The file has no readable data.",
        )

    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail=f"Could not process file: {str(error)}",
        )


# --------------------------------------------------
# DATA CLEANING
# --------------------------------------------------

@app.post("/api/clean/{dataset_id}/preview")
def preview_cleaning(dataset_id: str, payload: dict):
    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your CSV again.",
        )

    operations = payload.get("operations", {})
    cleaned_df, changes = clean_dataframe(
        DATASETS[dataset_id]["df"],
        operations,
    )
    return {
        "profile": profile(cleaned_df),
        "quality": data_quality_summary(cleaned_df),
        "preview": dataframe_records(cleaned_df.head(10)),
        "changes": changes,
    }


@app.post("/api/clean/{dataset_id}/apply")
def apply_cleaning(dataset_id: str, payload: dict):
    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your CSV again.",
        )

    cleaned_df, changes = clean_dataframe(
        DATASETS[dataset_id]["df"],
        payload.get("operations", {}),
    )
    DATASETS[dataset_id]["df"] = cleaned_df

    return {
        "success": True,
        "profile": profile(cleaned_df),
        "quality": data_quality_summary(cleaned_df),
        "preview": dataframe_records(cleaned_df.head(10)),
        "changes": changes,
    }


@app.post("/api/clean/{dataset_id}/reset")
def reset_cleaning(dataset_id: str):
    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your CSV again.",
        )

    original_df = DATASETS[dataset_id]["original_df"]
    restored_df = original_df.copy(deep=True)
    DATASETS[dataset_id]["df"] = restored_df

    return {
        "success": True,
        "profile": profile(restored_df),
        "quality": data_quality_summary(restored_df),
        "preview": dataframe_records(restored_df.head(10)),
    }


@app.get("/api/clean/{dataset_id}/download")
def download_cleaned_csv(dataset_id: str):
    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your CSV again.",
        )

    csv_content = DATASETS[dataset_id]["df"].to_csv(index=False)
    return Response(
        content=csv_content,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": 'attachment; filename="cleaned_dataset.csv"',
        },
    )


# --------------------------------------------------
# CHART DATA
# --------------------------------------------------

@app.get("/api/chart/{dataset_id}")
def get_chart_data(dataset_id: str):
    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your CSV again.",
        )

    df = DATASETS[dataset_id]["df"]

    numeric_columns = list(
        df.select_dtypes(include=np.number).columns
    )

    categorical_columns = list(
        df.select_dtypes(
            include=["object", "category", "bool"]
        ).columns
    )

    charts = []

    # First chart: numeric measure by the first categorical column.
    # Date-like categories are sorted chronologically.
    if categorical_columns and numeric_columns:
        category = categorical_columns[0]
        numeric = numeric_columns[0]

        summary = group_summary(
            df,
            category,
            numeric,
            operation="sum",
        )

        summary = sort_chart_data(summary, category)

        charts.append({
            "title": f"{numeric} by {category}",
            "type": "bar",
            "x_key": str(category),
            "y_key": str(numeric),
            "data": dataframe_records(summary),
        })

    # Additional charts: one for every remaining categorical column.
    # For sales.csv this adds Sales by Region and Sales by Product.
    if numeric_columns:
        numeric = numeric_columns[0]

        for category in categorical_columns[1:]:
            summary = group_summary(
                df,
                category,
                numeric,
                operation="sum",
            )

            charts.append({
                "title": f"{numeric} by {category}",
                "type": "bar",
                "x_key": str(category),
                "y_key": str(numeric),
                "data": dataframe_records(summary),
            })

    # Numeric distribution histogram.
    if numeric_columns:
        numeric = numeric_columns[0]

        charts.append({
            "title": f"{numeric} distribution",
            "type": "histogram",
            "x_key": "value",
            "y_key": "count",
            "data": build_histogram(df[numeric]),
        })

    # Fallback for datasets without numeric columns.
    if not charts and len(df.columns) > 0:
        column = df.columns[0]

        counts = (
            df[column]
            .value_counts(dropna=False)
            .head(15)
            .rename_axis(str(column))
            .reset_index(name="count")
        )

        charts.append({
            "title": f"Count by {column}",
            "type": "bar",
            "x_key": str(column),
            "y_key": "count",
            "data": dataframe_records(counts),
        })

    return {
        "dataset_id": dataset_id,
        "charts": charts,
    }


# --------------------------------------------------
# DATASET ANALYTICS
# --------------------------------------------------

@app.get("/api/analytics/{dataset_id}")
def get_analytics(dataset_id: str):
    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your CSV again.",
        )

    df = DATASETS[dataset_id]["df"]
    numeric = df.select_dtypes(include=np.number)

    numeric_summary = {}

    for column in numeric.columns:
        series = numeric[column].dropna()

        numeric_summary[str(column)] = {
            "sum": clean_value(series.sum()),
            "mean": clean_value(series.mean()),
            "median": clean_value(series.median()),
            "min": clean_value(series.min()),
            "max": clean_value(series.max()),
            "std": clean_value(series.std()),
            "unique": int(series.nunique()),
        }

    categorical_summary = {}

    for column in df.select_dtypes(
        include=["object", "category", "bool"]
    ).columns:
        categorical_summary[str(column)] = [
            {
                "value": clean_value(value),
                "count": int(count),
            }
            for value, count in (
                df[column]
                .value_counts(dropna=False)
                .head(15)
                .items()
            )
        ]

    advanced_insights = build_advanced_insights(df)

    return {
        "dataset_id": dataset_id,
        "filename": DATASETS[dataset_id]["filename"],
        "profile": profile(df),
        "numeric_summary": numeric_summary,
        "categorical_summary": categorical_summary,
        "business_insights": build_business_insights(df),
        "advanced_insights": advanced_insights,
    }


@app.get("/api/report/excel/{dataset_id}")
def download_excel_report(dataset_id: str):
    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your file again.",
        )

    dataset = DATASETS[dataset_id]
    df = dataset["df"]
    dataset_profile = profile(df)
    business = build_business_insights(df)
    advanced = build_advanced_insights(df)
    workbook = Workbook()

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="686BEF",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    def write_table(sheet, headers, rows, missing_label="Unavailable"):
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        sheet.freeze_panes = "A2"

        for row in rows:
            sheet.append([
                missing_label if value is None else value
                for value in row
            ])

        for column_cells in sheet.columns:
            column_letter = column_cells[0].column_letter
            width = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            sheet.column_dimensions[column_letter].width = min(
                max(width + 2, 12),
                50,
            )

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

    summary = workbook.active
    summary.title = "Summary"
    write_table(
        summary,
        ["Metric", "Value"],
        [
            ["Filename", dataset["filename"]],
            ["Dataset ID", dataset_id],
            ["Total Rows", dataset_profile["rows"]],
            ["Total Columns", dataset_profile["columns"]],
            ["Missing Cells", dataset_profile["missing_values"]],
            ["Duplicate Rows", dataset_profile["duplicate_rows"]],
            ["Total Sales", business["total_sales"]],
            ["Average Sales", business["average_sales"]],
            ["Average Quantity", advanced["average_quantity"]],
        ],
    )

    business_sheet = workbook.create_sheet("Business Insights")
    top_product = business["top_product"]
    lowest_product = business["lowest_product"]
    best_region = business["best_region"]
    highest_sale = business["highest_sale"]
    business_rows = [
        [
            "Top-performing product",
            (
                f"{top_product['name']} — {fmt(top_product['sales'])} in sales"
                if top_product
                else "Unavailable"
            ),
        ],
        [
            "Lowest-performing product",
            (
                f"{lowest_product['name']} — "
                f"{fmt(lowest_product['sales'])} in sales"
                if lowest_product
                else "Unavailable"
            ),
        ],
        [
            "Strongest region",
            (
                f"{best_region['name']} — {fmt(best_region['sales'])} in sales"
                if best_region
                else "Unavailable"
            ),
        ],
        [
            "Highest individual sale",
            fmt(highest_sale) if highest_sale is not None else "Unavailable",
        ],
        [
            "Overall sales performance",
            (
                f"Total sales: {fmt(business['total_sales'])}; "
                f"average sales: {fmt(business['average_sales'])}"
                if business["total_sales"] is not None
                or business["average_sales"] is not None
                else "Unavailable"
            ),
        ],
    ]
    recommendations = business["recommendations"]
    if recommendations:
        business_rows.extend(
            [
                [f"Recommendation {index}", recommendation]
                for index, recommendation in enumerate(
                    recommendations,
                    start=1,
                )
            ]
        )
    else:
        business_rows.append(["Recommendations", "Unavailable"])
    write_table(
        business_sheet,
        ["Insight", "Value"],
        business_rows,
    )

    numeric = df.select_dtypes(include=np.number)
    statistics_sheet = workbook.create_sheet("Statistics")
    write_table(
        statistics_sheet,
        ["Column", "Sum", "Average", "Median", "Minimum", "Maximum"],
        [
            [
                str(column),
                clean_value(series.sum()),
                clean_value(series.mean()),
                clean_value(series.median()),
                clean_value(series.min()),
                clean_value(series.max()),
            ]
            for column in numeric.columns
            for series in [numeric[column].dropna()]
        ],
    )

    quality_sheet = workbook.create_sheet("Data Quality")
    write_table(
        quality_sheet,
        ["Column", "Data Type", "Missing Values", "Unique Values"],
        [
            [
                str(column),
                str(df[column].dtype),
                int(df[column].isna().sum()),
                int(df[column].nunique(dropna=True)),
            ]
            for column in df.columns
        ],
    )

    raw_data_sheet = workbook.create_sheet("Raw Data")
    write_table(
        raw_data_sheet,
        [str(column) for column in df.columns],
        [
            [
                clean_value(value)
                for value in row
            ]
            for row in df.itertuples(index=False, name=None)
        ],
        missing_label=None,
    )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="InsightAI_Report_{dataset_id}.xlsx"'
            ),
        },
    )


@app.post("/api/compare")
def compare_datasets(payload: dict):
    dataset_1_id = payload.get("dataset_1_id")
    dataset_2_id = payload.get("dataset_2_id")

    if not isinstance(dataset_1_id, str) or dataset_1_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found: dataset_1_id.",
        )

    if not isinstance(dataset_2_id, str) or dataset_2_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found: dataset_2_id.",
        )

    def dataset_metrics(dataset_id):
        dataset = DATASETS[dataset_id]
        df = dataset["df"]
        metrics = {
            "dataset_id": dataset_id,
            "filename": dataset["filename"],
            "rows": int(df.shape[0]),
            "columns": int(df.shape[1]),
            "missing_cells": int(df.isna().sum().sum()),
            "duplicate_rows": int(df.duplicated().sum()),
        }

        for column, total_key, average_key in (
            ("Sales", "total_sales", "average_sales"),
            ("Quantity", "total_quantity", "average_quantity"),
        ):
            if (
                column in df.columns
                and pd.api.types.is_numeric_dtype(df[column])
            ):
                values = pd.to_numeric(
                    df[column],
                    errors="coerce",
                ).replace([np.inf, -np.inf], np.nan).dropna()
                total = values.sum(min_count=1)
                average = values.mean()
                metrics[total_key] = clean_value(total)
                metrics[average_key] = (
                    round(float(average), 2)
                    if pd.notna(average) and np.isfinite(average)
                    else None
                )
            else:
                metrics[total_key] = None
                metrics[average_key] = None

        return metrics

    dataset_1 = dataset_metrics(dataset_1_id)
    dataset_2 = dataset_metrics(dataset_2_id)
    changes = {}

    for metric in (
        "rows",
        "columns",
        "missing_cells",
        "duplicate_rows",
        "total_sales",
        "average_sales",
        "total_quantity",
        "average_quantity",
    ):
        value_1 = dataset_1[metric]
        value_2 = dataset_2[metric]

        if value_1 is None or value_2 is None:
            absolute_change = None
            percentage_change = None
        else:
            absolute_change = clean_value(value_2 - value_1)
            percentage_change = (
                round((value_2 - value_1) / value_1 * 100, 2)
                if value_1 != 0
                else None
            )

        changes[metric] = {
            "absolute_change": absolute_change,
            "percentage_change": percentage_change,
        }

    return {
        "dataset_1": dataset_1,
        "dataset_2": dataset_2,
        "changes": changes,
    }


# --------------------------------------------------
# AI ANALYST
# --------------------------------------------------

@app.post("/api/ask")
async def ask_question(payload: dict):
    dataset_id = payload.get("dataset_id")
    question = str(payload.get("question", "")).strip()

    if not dataset_id:
        raise HTTPException(
            status_code=400,
            detail="Please upload a dataset first.",
        )

    if dataset_id not in DATASETS:
        raise HTTPException(
            status_code=404,
            detail="Dataset not found. Upload your CSV again.",
        )

    if not question:
        raise HTTPException(
            status_code=400,
            detail="Please enter a question.",
        )

    if len(question) > 1000:
        raise HTTPException(
            status_code=400,
            detail="Question is too long. Keep it under 1000 characters.",
        )

    df = DATASETS[dataset_id]["df"]
    q = question.lower()

    numeric_columns = list(
        df.select_dtypes(include=np.number).columns
    )

    categorical_columns = list(
        df.select_dtypes(
            include=["object", "category", "bool"]
        ).columns
    )

    def response(answer, mode="calculated"):
        return {
            "question": question,
            "answer": str(answer),
            "mode": mode,
        }

    def numeric_column_from_question():
        column = find_mentioned_column(
            df,
            question,
            numeric_columns,
        )

        if column is not None:
            return column

        if numeric_columns and contains_any(q, [
            "average",
            "mean",
            "median",
            "standard deviation",
            "variance",
            "sum",
            "total",
            "highest",
            "maximum",
            "lowest",
            "minimum",
            "largest",
            "smallest",
        ]):
            return numeric_columns[0]

        return None

    # Dataset dimensions.
    if contains_any(q, [
        "how many rows",
        "number of rows",
        "row count",
        "total rows",
        "how many records",
        "number of records",
        "how many entries",
    ]):
        return response(
            f"The dataset contains {len(df):,} rows."
        )

    if contains_any(q, [
        "how many columns",
        "number of columns",
        "column count",
        "total columns",
        "how many features",
    ]):
        return response(
            f"The dataset contains {len(df.columns):,} columns: "
            + ", ".join(map(str, df.columns))
            + "."
        )

    # Missing values.
    if contains_any(q, [
        "missing values",
        "missing data",
        "null values",
        "null data",
        "isnull",
        "empty cells",
        "missing cells",
    ]):
        missing = df.isna().sum()
        total = int(missing.sum())

        details = [
            f"{column}: {int(count)}"
            for column, count in missing.items()
            if count > 0
        ]

        if total == 0:
            answer = "There are no missing values in this dataset."
        else:
            answer = (
                f"There are {total:,} missing cells. "
                + "; ".join(details)
                + "."
            )

        return response(answer)

    # Duplicate rows.
    if contains_any(q, [
        "duplicate",
        "duplicates",
        "repeated rows",
        "duplicate records",
    ]):
        duplicates = int(df.duplicated().sum())

        if duplicates == 0:
            answer = "There are no duplicate rows in this dataset."
        else:
            answer = (
                f"The dataset contains {duplicates:,} duplicate rows."
            )

        return response(answer)

    # Column names.
    if contains_any(q, [
        "list columns",
        "column names",
        "what are the columns",
        "show columns",
        "available columns",
    ]):
        return response(
            "The columns are: "
            + ", ".join(map(str, df.columns))
            + "."
        )

    # Data types.
    if contains_any(q, [
        "data types",
        "column types",
        "dtypes",
        "type of each column",
    ]):
        details = [
            f"{column}: {dtype}"
            for column, dtype in df.dtypes.items()
        ]

        return response(
            "Column data types: " + "; ".join(details) + "."
        )

    # Unique values.
    if contains_any(q, [
        "unique values",
        "distinct values",
        "how many unique",
        "how many distinct",
        "number of unique",
    ]):
        column = find_mentioned_column(df, question)

        if column is not None:
            unique_count = int(df[column].nunique(dropna=True))
            values = df[column].dropna().unique().tolist()

            if len(values) <= 20:
                return response(
                    f"Column '{column}' has {unique_count:,} unique "
                    f"values: {', '.join(map(str, values))}."
                )

            return response(
                f"Column '{column}' has {unique_count:,} unique values."
            )

    # Category comparison and rankings.
    group_column = find_mentioned_column(
        df,
        question,
        categorical_columns,
    )

    value_column = numeric_column_from_question()

    ranking_question = contains_any(q, [
        "highest",
        "lowest",
        "largest",
        "smallest",
        "maximum",
        "minimum",
        "most",
        "least",
        "top",
        "bottom",
        "best",
        "worst",
    ])

    average_question = contains_any(q, [
        "average",
        "mean",
        "avg",
    ])

    comparison_question = wants_group_comparison(q)

    if group_column is not None and (
        comparison_question
        or ranking_question
        or "by " in q
        or "per " in q
    ):
        if value_column is None and numeric_columns:
            value_column = numeric_columns[0]

        if value_column is not None:
            operation = "mean" if average_question else "sum"

            summary = group_summary(
                df,
                group_column,
                value_column,
                operation=operation,
            )

            if summary.empty:
                return response(
                    f"No usable numeric values were found in '{value_column}'."
                )

            if ranking_question:
                highest = contains_any(q, [
                    "highest",
                    "largest",
                    "maximum",
                    "most",
                    "top",
                    "best",
                ])

                lowest = contains_any(q, [
                    "lowest",
                    "smallest",
                    "minimum",
                    "least",
                    "bottom",
                    "worst",
                ])

                if lowest and not highest:
                    selected = summary.iloc[-1]
                    label = "lowest"
                else:
                    selected = summary.iloc[0]
                    label = "highest"

                category_value = selected[group_column]
                metric_value = selected[value_column]

                answer = (
                    f"{category_value} has the {label} "
                    f"{operation} {value_column}: {fmt(metric_value)}."
                )

                if contains_any(q, [
                    "compare",
                    "breakdown",
                    "each",
                    "every",
                    "all",
                ]):
                    details = [
                        f"{row[group_column]}: {fmt(row[value_column])}"
                        for _, row in summary.iterrows()
                    ]

                    answer += (
                        " Breakdown: " + "; ".join(details) + "."
                    )

                return response(answer)

            details = [
                f"{row[group_column]}: {fmt(row[value_column])}"
                for _, row in summary.iterrows()
            ]

            operation_label = (
                "average" if operation == "mean" else "total"
            )

            return response(
                f"{operation_label.title()} {value_column} by "
                f"{group_column}: "
                + "; ".join(details)
                + "."
            )

    # Numeric statistics.
    value_column = numeric_column_from_question()

    if value_column is not None:
        series = pd.to_numeric(
            df[value_column],
            errors="coerce",
        ).dropna()

        if series.empty:
            return response(
                f"Column '{value_column}' contains no usable numeric values."
            )

        if contains_any(q, [
            "standard deviation",
            "std deviation",
            "spread of",
        ]):
            return response(
                f"The standard deviation of {value_column} is "
                f"{fmt(series.std())}."
            )

        if "variance" in q:
            return response(
                f"The variance of {value_column} is "
                f"{fmt(series.var())}."
            )

        if "median" in q:
            return response(
                f"The median of {value_column} is "
                f"{fmt(series.median())}."
            )

        if contains_any(q, ["average", "mean", "avg"]):
            return response(
                f"The average {value_column} is "
                f"{fmt(series.mean())}."
            )

        if contains_any(q, ["sum", "total", "overall"]):
            return response(
                f"The total {value_column} is "
                f"{fmt(series.sum())}."
            )

        if contains_any(q, [
            "highest",
            "maximum",
            "largest",
            "maximum value",
        ]):
            return response(
                f"The highest {value_column} is "
                f"{fmt(series.max())}."
            )

        if contains_any(q, [
            "lowest",
            "minimum",
            "smallest",
            "minimum value",
        ]):
            return response(
                f"The lowest {value_column} is "
                f"{fmt(series.min())}."
            )

        if contains_any(q, [
            "summary statistics",
            "describe",
            "statistics of",
            "statistical summary",
        ]):
            return response(
                f"Statistics for {value_column}: "
                f"count={fmt(series.count())}, "
                f"sum={fmt(series.sum())}, "
                f"average={fmt(series.mean())}, "
                f"median={fmt(series.median())}, "
                f"minimum={fmt(series.min())}, "
                f"maximum={fmt(series.max())}, "
                f"standard deviation={fmt(series.std())}."
            )

    # Dataset summary.
    if contains_any(q, [
        "summarize",
        "summary",
        "describe dataset",
        "overview",
        "analyze this dataset",
        "tell me about this dataset",
        "what can you tell me",
    ]):
        numeric_summary = []

        for column in numeric_columns:
            series = pd.to_numeric(
                df[column],
                errors="coerce",
            ).dropna()

            if not series.empty:
                numeric_summary.append(
                    f"{column}: total={fmt(series.sum())}, "
                    f"average={fmt(series.mean())}, "
                    f"minimum={fmt(series.min())}, "
                    f"maximum={fmt(series.max())}"
                )

        categorical_summary = []

        for column in categorical_columns:
            categorical_summary.append(
                f"{column} has "
                f"{df[column].nunique(dropna=True)} unique values"
            )

        parts = [
            f"The dataset has {len(df):,} rows and "
            f"{len(df.columns):,} columns.",
            f"There are {int(df.isna().sum().sum())} missing cells.",
            f"There are {int(df.duplicated().sum())} duplicate rows.",
        ]

        if numeric_summary:
            parts.append(
                "Numeric summary: " + "; ".join(numeric_summary) + "."
            )

        if categorical_summary:
            parts.append(
                "Categorical summary: "
                + "; ".join(categorical_summary)
                + "."
            )

        return response(" ".join(parts))

    # Groq AI fallback for open-ended questions.
    client = get_groq_client()

    if client is None:
        return response(
            "I can answer questions about row counts, columns, missing "
            "values, duplicates, unique values, numeric statistics, "
            "category rankings, and dataset summaries. For more open-ended "
            "AI analysis, add GROQ_API_KEY to your backend .env file.",
            mode="fallback",
        )

    try:
        dataset_context = build_dataset_context(df)

        completion = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are InsightAI, a careful data analyst. "
                        "Answer questions using only the supplied dataset "
                        "context. Never invent values or claim to have "
                        "computed something not supported by the data. "
                        "Be clear, concise, and explain relevant patterns. "
                        "If the context is insufficient, say so."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"Dataset context:\n{dataset_context}\n\n"
                        f"User question:\n{question}\n\n"
                        "Provide a direct answer. Use readable numbers "
                        "and mention any important limitations."
                    ),
                },
            ],
            temperature=0.2,
            max_tokens=800,
        )

        answer = completion.choices[0].message.content

        if not answer:
            answer = (
                "The AI could not generate an answer. Please try again."
            )

        return response(answer, mode="ai")

    except Exception as error:
        print(f"Groq API error: {error}")

        return response(
            "The AI service could not answer this question right now. "
            "Try asking for a specific statistic, such as the average "
            "of a numeric column or the highest-selling product. "
            "If this continues, check your GROQ_API_KEY and model setting.",
            mode="fallback",
        )